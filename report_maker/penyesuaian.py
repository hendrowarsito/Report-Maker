"""
Penyesuaian laporan dengan AI tanpa placeholder.

Claude membaca dokumen sumber (PDF sertifikat, lembar kerja Excel) dan laporan
yang sudah terisi, lalu MENGUSULKAN penggantian kalimat yang tidak sesuai data.
Tidak ada yang ditulis ke laporan sebelum penilai menyetujui usulan.
"""
import base64
import json
from dataclasses import dataclass
from io import BytesIO

import anthropic
import pandas as pd
from docx import Document
from openpyxl.utils import get_column_letter

from .docx_engine import _W_T, _replace_in_para_xml, iter_all_paragraphs
from .terbilang import format_nilai

MODEL_PENYESUAIAN = "claude-opus-5-5"


# =============================================================================
# Paragraf laporan — diberi nomor agar usulan Claude bisa ditempatkan tepat
# =============================================================================

def _paragraf_unik(doc: Document):
    """Paragraf dokumen tanpa duplikat (sel gabungan & header bertaut muncul dua kali)."""
    dilihat = set()
    for para in iter_all_paragraphs(doc):
        if para._element in dilihat:
            continue
        dilihat.add(para._element)
        yield para


def _teks_paragraf(para) -> str:
    return "".join(t.text or "" for t in para._element.iter(_W_T))


def daftar_paragraf(docx_bytes: bytes) -> dict[int, str]:
    """{nomor: teks} untuk setiap paragraf yang berisi teks. Nomor stabil per dokumen."""
    doc = Document(BytesIO(docx_bytes))
    return {
        i: teks
        for i, para in enumerate(_paragraf_unik(doc))
        if (teks := _teks_paragraf(para)).strip()
    }


# =============================================================================
# Dokumen sumber
# =============================================================================

def excel_ke_teks(xlsx_bytes: bytes) -> str:
    """
    Semua sheet Excel sebagai teks berlabel sel (mis. "B12: 312"), agar Claude
    bisa menyebut sumbernya. Nilai rumus diambil dari hasil terakhir yang
    disimpan Excel.
    """
    bagian = []
    for nama_sheet, df in pd.read_excel(
        BytesIO(xlsx_bytes), sheet_name=None, header=None, engine="openpyxl"
    ).items():
        baris_teks = []
        for r, row in enumerate(df.itertuples(index=False), start=1):
            sel = [
                f"{get_column_letter(c)}{r}: {teks}"
                for c, nilai in enumerate(row, start=1)
                if (teks := format_nilai(nilai))
            ]
            if sel:
                baris_teks.append(" | ".join(sel))
        if baris_teks:
            bagian.append(f"## Sheet: {nama_sheet}\n" + "\n".join(baris_teks))
    return "\n\n".join(bagian)


# =============================================================================
# Usulan dari Claude
# =============================================================================

@dataclass
class Usulan:
    paragraf: int
    teks_lama: str
    teks_baru: str
    alasan: str
    sumber: str
    kutipan_sumber: str


_SKEMA = {
    "type": "object",
    "properties": {
        "usulan": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "paragraf": {"type": "integer"},
                    "teks_lama": {"type": "string"},
                    "teks_baru": {"type": "string"},
                    "alasan": {"type": "string"},
                    "sumber": {"type": "string"},
                    "kutipan_sumber": {"type": "string"},
                },
                "required": [
                    "paragraf", "teks_lama", "teks_baru",
                    "alasan", "sumber", "kutipan_sumber",
                ],
                "additionalProperties": False,
            },
        },
        "catatan": {"type": "string"},
    },
    "required": ["usulan", "catatan"],
    "additionalProperties": False,
}

_SYSTEM = """\
Anda membantu penilai di Kantor Jasa Penilai Publik (KJPP) memeriksa laporan \
penilaian properti sebelum diterbitkan. Anda menerima dokumen sumber (sertifikat \
tanah, lembar kerja penilaian, dan dokumen pendukung lain) serta isi laporan yang \
paragrafnya diberi nomor [n].

Tugas: temukan bagian laporan yang bertentangan dengan dokumen sumber atau masih \
memuat data objek lain dari template lama — misalnya nomor dan jenis sertifikat, \
NIB, luas tanah/bangunan, nama pemegang hak, alamat dan batas lokasi, tanggal \
(inspeksi, penilaian, laporan, terbit sertifikat), nilai, NJOP, dan terbilangnya — \
lalu usulkan penggantinya.

Aturan:
- Usulkan perubahan hanya bila dokumen sumber jelas mendukungnya. Jangan menebak. \
Bila sumber buram, saling bertentangan, atau tidak lengkap, jangan buat usulan; \
jelaskan di "catatan".
- "teks_lama" harus disalin persis dari teks paragraf (huruf, spasi, tanda baca), \
sependek mungkin tetapi tetap unik di paragraf itu — cukup frasa atau kalimat \
yang berubah, bukan seluruh paragraf.
- "teks_baru" mengikuti gaya bahasa dan format laporan (format angka Indonesia, \
huruf kapital, satuan). Bila angka berubah, sesuaikan juga terbilangnya dalam \
usulan yang sama.
- Jangan mengubah kalimat baku, metodologi, asumsi, atau opini penilai yang tidak \
terkait data objek.
- "sumber" menyebut nama file dan lokasinya (halaman PDF, atau sheet dan sel Excel). \
"kutipan_sumber" menyalin teks pendukung apa adanya dari sumber.
- "alasan" dan "catatan" ditulis singkat dalam bahasa Indonesia.
- Bila tidak ada yang perlu diubah, kembalikan daftar usulan kosong.\
"""


def usulkan_penyesuaian(
    client: anthropic.Anthropic,
    paragraf: dict[int, str],
    pdfs: list[tuple[str, bytes]],
    excels: list[tuple[str, str]],
) -> tuple[list[Usulan], str]:
    """
    Minta Claude membandingkan laporan dengan dokumen sumber.
    pdfs: [(nama file, isi PDF)]; excels: [(nama file, hasil excel_ke_teks)].
    Kembalikan (usulan, catatan). RuntimeError bila permintaan ditolak/terpotong.
    """
    konten: list[dict] = []
    for nama, isi in pdfs:
        konten.append({
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": "application/pdf",
                "data": base64.standard_b64encode(isi).decode(),
            },
            "title": nama,
        })
    for nama, teks in excels:
        konten.append({"type": "text", "text": f"# Lembar kerja Excel: {nama}\n\n{teks}"})
    if konten:
        # Sumber jarang berubah antar-percobaan: cache agar analisis ulang lebih murah
        konten[-1]["cache_control"] = {"type": "ephemeral"}

    laporan = "\n".join(f"[{i}] {teks}" for i, teks in paragraf.items())
    konten.append({
        "type": "text",
        "text": f"# Isi laporan (paragraf bernomor)\n\n{laporan}\n\n"
                "Bandingkan laporan dengan dokumen sumber di atas dan berikan usulan.",
    })

    with client.beta.messages.stream(
        model=MODEL_PENYESUAIAN,
        max_tokens=64000,
        system=_SYSTEM,
        messages=[{"role": "user", "content": konten}],
        output_config={
            "effort": "high",
            "format": {"type": "json_schema", "schema": _SKEMA},
        },
        # Bila ditolak filter keamanan, ulangi otomatis di model cadangan Anthropic
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
    ) as stream:
        respons = stream.get_final_message()

    if respons.stop_reason == "refusal":
        raise RuntimeError("Permintaan ditolak oleh Claude. Coba lagi atau periksa dokumen sumber.")
    if respons.stop_reason == "max_tokens":
        raise RuntimeError("Jawaban Claude terpotong. Coba dengan dokumen sumber yang lebih sedikit.")

    teks = next(b.text for b in respons.content if b.type == "text")
    data = json.loads(teks)
    return [Usulan(**u) for u in data["usulan"]], data["catatan"]


def validasi_usulan(
    paragraf: dict[int, str], usulan: list[Usulan]
) -> tuple[list[Usulan], list[tuple[Usulan, str]]]:
    """Pisahkan usulan yang bisa diterapkan dari yang tidak (beserta alasannya)."""
    valid, tidak_valid = [], []
    for u in usulan:
        if u.paragraf not in paragraf:
            tidak_valid.append((u, f"Paragraf [{u.paragraf}] tidak ada"))
        elif not u.teks_lama or u.teks_lama not in paragraf[u.paragraf]:
            tidak_valid.append((u, "Teks lama tidak ditemukan di paragraf"))
        elif u.teks_lama == u.teks_baru:
            tidak_valid.append((u, "Tidak ada perubahan"))
        else:
            valid.append(u)
    return valid, tidak_valid


def terapkan_penyesuaian(docx_bytes: bytes, usulan: list[Usulan]) -> tuple[bytes, int]:
    """Terapkan usulan yang disetujui; format teks di sekitarnya tetap terjaga."""
    per_paragraf: dict[int, list[Usulan]] = {}
    for u in usulan:
        per_paragraf.setdefault(u.paragraf, []).append(u)

    doc = Document(BytesIO(docx_bytes))
    jumlah = 0
    for i, para in enumerate(_paragraf_unik(doc)):
        for u in per_paragraf.get(i, []):
            jumlah += _replace_in_para_xml(para, u.teks_lama, u.teks_baru)

    out = BytesIO()
    doc.save(out)
    return out.getvalue(), jumlah
