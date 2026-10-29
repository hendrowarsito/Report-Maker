"""Pengisian placeholder dan Find & Replace di dokumen Word (.docx)."""
import re
from io import BytesIO
from pathlib import Path

from docx import Document


# =============================================================================
# Paragraph-level placeholder replacement — berbasis XML langsung
# =============================================================================
#
# MENGAPA XML, BUKAN para.runs?
# para.runs hanya mengembalikan <w:r> anak-langsung dari <w:p>.
# Placeholder bisa berada di dalam elemen pembungkus yang TIDAK terlihat oleh
# para.runs, antara lain:
#   • <w:hyperlink>   — teks yang merupakan hyperlink
#   • <w:ins>         — teks hasil tracked-change (insertion)
#   • <w:sdt>         — content control (structured document tag)
# Dengan para._element.iter(W_T) kita mendapatkan SEMUA <w:t> di seluruh
# sub-pohon XML, sehingga tidak ada placeholder yang terlewat.
# Karena kita hanya mengubah .text pada elemen <w:t> (bukan <w:rPr>),
# semua formatting (bold/italic/warna/font) tetap terjaga otomatis.
#
_W_NS  = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_T   = f"{{{_W_NS}}}t"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _set_t_text(t_elem, text: str):
    """Tulis teks ke elemen <w:t>, pasang xml:space="preserve" jika ada spasi di tepi."""
    t_elem.text = text
    if text and (text[0] == " " or text[-1] == " "):
        t_elem.set(_XML_SPACE, "preserve")


def _replace_in_para_xml(para, placeholder: str, value: str) -> int:
    """
    Ganti semua kemunculan placeholder di paragraf secara langsung via XML.
    Menangani: run normal, hyperlink, tracked-change, content-control,
               dan placeholder yang terpecah di beberapa elemen <w:t>.
    Kembalikan jumlah penggantian yang dilakukan.
    """
    t_elems = list(para._element.iter(_W_T))
    if not t_elems:
        return 0

    count = 0

    # Babak 1: ganti dalam satu elemen <w:t> (kasus paling umum)
    for t in t_elems:
        txt = t.text or ""
        if placeholder in txt:
            count += txt.count(placeholder)
            _set_t_text(t, txt.replace(placeholder, value))

    # Babak 2: ganti placeholder yang terpecah di beberapa <w:t>
    # Loop sampai tidak ada lagi kemunculan (menangani beberapa kejadian sekaligus)
    for _ in range(len(t_elems)):           # batas atas iterasi yang aman
        texts    = [t.text or "" for t in t_elems]
        full     = "".join(texts)
        if placeholder not in full:
            break

        ph_start = full.index(placeholder)
        ph_end   = ph_start + len(placeholder)

        # Bangun posisi kumulatif setiap elemen
        cum, cum_pos = 0, []
        for txt in texts:
            cum_pos.append(cum)
            cum += len(txt)

        s_idx = e_idx = None
        for i, (cp, txt) in enumerate(zip(cum_pos, texts)):
            cp_end = cp + len(txt)
            if cp <= ph_start < cp_end:
                s_idx = i
            if cp < ph_end <= cp_end:
                e_idx = i

        # Tangani placeholder yang tepat berakhir di batas elemen
        if e_idx is None:
            for i, cp in enumerate(cum_pos):
                if cp + len(texts[i]) == ph_end:
                    e_idx = i
                    break

        if s_idx is None or e_idx is None or s_idx == e_idx:
            break   # tidak dapat diselesaikan, hentikan agar tidak infinite loop

        before = texts[s_idx][: ph_start - cum_pos[s_idx]]
        after  = texts[e_idx][ph_end   - cum_pos[e_idx] :]

        _set_t_text(t_elems[s_idx], before + value)
        _set_t_text(t_elems[e_idx], after)
        for i in range(s_idx + 1, e_idx):
            t_elems[i].text = ""

        count += 1

    return count


def replace_placeholders_in_paragraph(para, data: dict):
    """
    Satu lintasan: cari placeholder yang benar-benar ada di paragraf, lalu ganti
    hanya yang punya nilai — tidak mengulang seluruh key data untuk tiap paragraf.
    """
    teks = "".join(t.text or "" for t in para._element.iter(_W_T))
    if "{{" not in teks:
        return para
    for nama in set(_POLA_PLACEHOLDER.findall(teks)):
        if nama in data:
            _replace_in_para_xml(para, f"{{{{{nama}}}}}", str(data[nama]))
    return para


# =============================================================================
# Document-level replacement (body + tabel tersarang + header + footer)
# =============================================================================

def _iter_table_paragraphs(tables):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                yield from _iter_table_paragraphs(cell.tables)


def iter_all_paragraphs(doc: Document):
    """Semua paragraf dokumen: body, tabel tersarang, header, dan footer."""
    yield from doc.paragraphs
    yield from _iter_table_paragraphs(doc.tables)

    for section in doc.sections:
        for hf in (
            section.header, section.footer,
            section.even_page_header, section.even_page_footer,
            section.first_page_header, section.first_page_footer,
        ):
            if hf is not None:
                yield from hf.paragraphs
                yield from _iter_table_paragraphs(hf.tables)


def replace_placeholders(doc: Document, data: dict) -> Document:
    for para in iter_all_paragraphs(doc):
        replace_placeholders_in_paragraph(para, data)
    return doc


_POLA_PLACEHOLDER = re.compile(r"\{\{([^{}]+)\}\}")


def collect_placeholders(doc: Document) -> set[str]:
    """Nama semua placeholder {{...}} di dokumen, termasuk header/footer."""
    nama: set[str] = set()
    for para in iter_all_paragraphs(doc):
        teks = "".join(t.text or "" for t in para._element.iter(_W_T))
        nama.update(_POLA_PLACEHOLDER.findall(teks))
    return nama


# =============================================================================
# Find & Replace langsung (tanpa placeholder, tanpa API)
# =============================================================================

def apply_find_replace(doc: Document, pairs: list[tuple[str, str]]) -> int:
    """
    Ganti kata secara langsung di seluruh dokumen (body, tabel, header, footer).
    Menggunakan _replace_in_para_xml yang sama dengan placeholder replacement
    sehingga hyperlink, tracked-change, dan split antar <w:t> ikut tertangani.
    Kembalikan jumlah total penggantian yang dilakukan.
    """
    total = 0
    substitutions = [(f, r) for f, r in pairs if f.strip()]

    for para in iter_all_paragraphs(doc):
        for find, replace in substitutions:
            total += _replace_in_para_xml(para, find, replace)

    return total


# =============================================================================
# File processing helpers
# =============================================================================

def placeholders_dari_bytes(docx_bytes: bytes) -> set[str]:
    return collect_placeholders(Document(BytesIO(docx_bytes)))


def isi_dokumen(
    docx_bytes: bytes,
    data: dict[str, str],
    find_replace_pairs: list[tuple[str, str]] | tuple = (),
) -> tuple[bytes, str, list[str]]:
    """
    Isi template dengan data lalu terapkan Find & Replace.
    Kembalikan (docx hasil, teks untuk preview/cek, placeholder yang tersisa).
    """
    doc = Document(BytesIO(docx_bytes))
    replace_placeholders(doc, data)
    if find_replace_pairs:
        apply_find_replace(doc, list(find_replace_pairs))
    output = BytesIO()
    doc.save(output)
    sisa = sorted(collect_placeholders(doc), key=str.lower)
    return output.getvalue(), extract_text_from_docx(doc), sisa


def extract_text_from_docx(docx_obj: Document) -> str:
    """
    Ekstrak semua teks: paragraf tubuh dokumen + isi sel tabel.
    Versi lama hanya mengambil doc.paragraphs sehingga tabel tidak ikut dipreview.
    """
    parts: list[str] = []

    for para in docx_obj.paragraphs:
        parts.append(para.text)

    def _collect_table(table):
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                cell_text = " | ".join(
                    para.text for para in cell.paragraphs if para.text.strip()
                )
                if cell_text:
                    row_cells.append(cell_text)
                for nested in cell.tables:
                    _collect_table(nested)
            if row_cells:
                parts.append("  ".join(row_cells))

    for table in docx_obj.tables:
        _collect_table(table)

    return "\n".join(parts)


def highlight_placeholders(text: str) -> str:
    return re.sub(r"(\{\{.*?\}\})", r"[PLACEHOLDER:\1]", text)


_KARAKTER_TERLARANG = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def nama_file_hasil(pola: str, nama_template: str, data: dict[str, str]) -> str:
    """
    Nama file hasil (tanpa ekstensi) dari pola berisi {{key}}.
    {{nama_file}} = nama template tanpa .docx; key yang tidak ada → kosong.
    Contoh: "LP_{{nama}}" + {"nama": "PT ABC"} → "LP_PT ABC"
    """
    stem = Path(nama_template).stem
    nilai = {**data, "nama_file": stem}
    hasil = _POLA_PLACEHOLDER.sub(lambda m: str(nilai.get(m.group(1), "")), pola)
    hasil = _KARAKTER_TERLARANG.sub("-", hasil)
    hasil = re.sub(r"\s+", " ", hasil).strip(" ._-")
    return hasil[:150] or stem
