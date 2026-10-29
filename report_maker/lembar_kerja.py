"""
Mengisi placeholder langsung dari lembar kerja asli — tanpa sheet "Laporan".

Claude memetakan setiap placeholder ke sel lembar kerja (mis. {{alamat}} →
"LK Tanah!C5"). Nilai selalu dibaca dari sel itu, bukan ditulis ulang AI.
Pemetaan bisa disimpan sebagai JSON dan dipakai lagi untuk laporan berikutnya
dengan format lembar kerja yang sama, tanpa memanggil AI.
"""
import datetime
import json
import numbers
import re
from dataclasses import asdict, dataclass, field
from io import BytesIO

import anthropic
import openpyxl

from .ai import minta_json
from .terbilang import format_nilai, pecah_turunan


# =============================================================================
# Membaca sel lembar kerja
# =============================================================================

def _angka_id(n: float, desimal: int) -> str:
    """1234567.5, 2 → '1.234.567,50'"""
    return f"{n:,.{desimal}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def format_sel(nilai, number_format: str | None = "General") -> str:
    """
    Teks sel seperti yang terlihat di Excel, bergaya Indonesia:
      0.25 dengan format 0%        → '25%'
      1500000 dengan format #,##0  → '1.500.000'
      lainnya                      → format_nilai (tanggal '28 April 2026', dst.)
    """
    fmt = number_format or "General"
    if isinstance(nilai, bool) or not isinstance(nilai, numbers.Real) \
            or isinstance(nilai, (datetime.date, datetime.time)):
        return format_nilai(nilai)
    desimal = len(m.group(1)) if (m := re.search(r"0\.(0+)", fmt)) else 0
    if "%" in fmt:
        return _angka_id(nilai * 100, desimal) + "%"
    if "#,##0" in fmt:
        return _angka_id(nilai, desimal)
    return format_nilai(nilai)


def baca_sel(xlsx_bytes: bytes) -> dict[str, str]:
    """
    {"Sheet!C5": teks} untuk semua sel berisi di semua sheet, berurutan.
    Nilai rumus diambil dari hasil terakhir yang disimpan Excel.
    """
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True, read_only=True)
    hasil: dict[str, str] = {}
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if teks := format_sel(cell.value, cell.number_format):
                        hasil[f"{ws.title}!{cell.coordinate}"] = teks
    finally:
        wb.close()
    return hasil


def sel_ke_teks(sel: dict[str, str]) -> str:
    """Sel dikelompokkan per sheet dan baris: '## Sheet: X\\nA1: Luas | B1: 312'."""
    per_sheet: dict[str, dict[int, list[str]]] = {}
    for ref, teks in sel.items():
        sheet, koord = ref.rsplit("!", 1)
        baris = int(re.sub(r"\D", "", koord))
        isi = teks.replace("\n", " / ")
        per_sheet.setdefault(sheet, {}).setdefault(baris, []).append(f"{koord}: {isi}")
    return "\n\n".join(
        f"## Sheet: {sheet}\n" + "\n".join(" | ".join(r) for r in baris.values())
        for sheet, baris in per_sheet.items()
    )


# =============================================================================
# Pemetaan placeholder → sel
# =============================================================================

KEYAKINAN = ["tinggi", "sedang", "rendah", "tidak_ditemukan"]


@dataclass
class Pemetaan:
    sel: list[str] = field(default_factory=list)
    pemisah: str = " "
    keyakinan: str = "tinggi"
    alasan: str = ""


def kebutuhan_kunci(placeholders: set[str]) -> dict[str, str]:
    """
    Key yang harus dicari di lembar kerja beserta jenisnya. Placeholder turunan
    ({{X_TERBILANG}}, {{X_FORMAT}}) tidak dicari sendiri: cukup X, sebagai angka.
    """
    kebutuhan: dict[str, str] = {}
    for ph in sorted(placeholders):
        sumber, _ = pecah_turunan(ph)
        if sumber:
            kebutuhan[sumber] = "angka"
        else:
            kebutuhan.setdefault(ph, "teks")
    return kebutuhan


def terapkan_pemetaan(
    pemetaan: dict[str, Pemetaan], sel: dict[str, str]
) -> tuple[dict[str, str], dict[str, str]]:
    """
    Baca nilai tiap key dari selnya. Kembalikan (data, peringatan).
    Key tanpa nilai tidak dimasukkan ke data, agar placeholdernya tetap
    terlihat di laporan alih-alih terhapus diam-diam.
    """
    data: dict[str, str] = {}
    peringatan: dict[str, str] = {}
    for key, p in pemetaan.items():
        hilang = [r for r in p.sel if r not in sel]
        nilai = p.pemisah.join(sel[r] for r in p.sel if r in sel)
        if nilai:
            data[key] = nilai
        if not p.sel:
            peringatan[key] = "Tidak ditemukan di lembar kerja"
        elif hilang:
            peringatan[key] = "Sel kosong/tidak ada: " + ", ".join(hilang)
        elif p.keyakinan == "rendah":
            peringatan[key] = "Keyakinan rendah — cek " + ", ".join(p.sel)
    return data, peringatan


_SKEMA = {
    "type": "object",
    "properties": {
        "pemetaan": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "key": {"type": "string"},
                    "sel": {"type": "array", "items": {"type": "string"}},
                    "pemisah": {"type": "string"},
                    "keyakinan": {"type": "string", "enum": KEYAKINAN},
                    "alasan": {"type": "string"},
                },
                "required": ["key", "sel", "pemisah", "keyakinan", "alasan"],
                "additionalProperties": False,
            },
        },
        "catatan": {"type": "string"},
    },
    "required": ["pemetaan", "catatan"],
    "additionalProperties": False,
}

_SYSTEM = """\
Anda membantu penilai di Kantor Jasa Penilai Publik (KJPP) mengisi template \
laporan penilaian properti dari lembar kerja Excel. Anda menerima isi lembar \
kerja per sel (format "Sheet!C5") dan daftar key placeholder beserta kalimat \
template tempat key itu dipakai.

Untuk setiap key, tentukan sel yang nilainya harus mengisi placeholder itu:
- Pahami arti key dari namanya DAN kalimat konteksnya, lalu cari sel nilai \
yang tepat — biasanya di sebelah atau di bawah sel label yang sesuai \
(mis. label "Alamat" di B5, nilainya di C5). Pilih sel berisi nilai, bukan label.
- Perhatikan sheet: data tanah ada di sheet tanah, data bangunan di sheet \
bangunan, nilai akhir di sheet rekapitulasi/kesimpulan, dst.
- Key berjenis "angka" harus menunjuk ke sel berisi angka (terbilangnya dibuat \
otomatis oleh aplikasi).
- Bila nilai tersusun dari beberapa sel (mis. jalan, kelurahan, kecamatan), \
cantumkan sel-selnya berurutan di "sel" dan isi "pemisah" (mis. ", "). Bila \
satu sel, "pemisah" diisi " ".
- Tulis referensi sel persis seperti di data ("Nama Sheet!C5").
- "keyakinan": tinggi bila jelas; sedang bila ada beberapa kandidat; rendah \
bila ragu; tidak_ditemukan (dengan "sel" kosong) bila tidak ada. Jangan menebak.
- "alasan" singkat dalam bahasa Indonesia, sebut label sel yang dipakai.
- Petakan setiap key yang diminta tepat satu kali.\
"""


def petakan_dengan_ai(
    client: anthropic.Anthropic,
    sel: dict[str, str],
    kebutuhan: dict[str, str],
    konteks: dict[str, str],
) -> tuple[dict[str, Pemetaan], str]:
    """
    Minta Claude memetakan key ke sel. konteks: {placeholder: kalimat template}.
    Kembalikan (pemetaan untuk key yang diminta, catatan).
    """
    def _konteks(key: str) -> str:
        if key in konteks:
            return konteks[key]
        return next((k for ph, k in konteks.items() if pecah_turunan(ph)[0] == key), "")

    daftar = "\n".join(
        f'- `{key}` ({jenis}) — konteks: "{_konteks(key)}"' for key, jenis in kebutuhan.items()
    )
    konten = [
        {
            "type": "text",
            "text": f"# Isi lembar kerja\n\n{sel_ke_teks(sel)}",
            # Lembar kerja yang sama dipakai ulang saat memetakan key berikutnya
            "cache_control": {"type": "ephemeral"},
        },
        {"type": "text", "text": f"# Key yang perlu dipetakan\n\n{daftar}"},
    ]
    data = minta_json(client, _SYSTEM, konten, _SKEMA)
    pemetaan = {
        p["key"]: Pemetaan(
            sel=p["sel"], pemisah=p["pemisah"], keyakinan=p["keyakinan"], alasan=p["alasan"]
        )
        for p in data["pemetaan"]
        if p["key"] in kebutuhan
    }
    return pemetaan, data["catatan"]


# =============================================================================
# Simpan / muat pemetaan
# =============================================================================

def pemetaan_ke_json(pemetaan: dict[str, Pemetaan]) -> str:
    return json.dumps(
        {"versi": 1, "pemetaan": {k: asdict(v) for k, v in sorted(pemetaan.items())}},
        ensure_ascii=False,
        indent=2,
    )


def pemetaan_dari_json(teks: str | bytes) -> dict[str, Pemetaan]:
    """ValueError bila file bukan pemetaan yang valid."""
    try:
        data = json.loads(teks)
        return {
            str(k): Pemetaan(
                sel=[str(r) for r in v.get("sel", [])],
                pemisah=str(v.get("pemisah", " ")),
                keyakinan=str(v.get("keyakinan", "tinggi")),
                alasan=str(v.get("alasan", "")),
            )
            for k, v in data["pemetaan"].items()
        }
    except (KeyError, TypeError, AttributeError, json.JSONDecodeError) as e:
        raise ValueError(f"File pemetaan tidak valid: {e}") from e
