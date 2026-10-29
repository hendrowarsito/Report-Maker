import shutil
from io import BytesIO

import pytest
from docx import Document

from report_maker.penyesuaian import Usulan
from report_maker.preview import (
    gambar_halaman,
    halaman_berubah,
    perubahan_html,
    teks_per_halaman,
)


def test_perubahan_html_coret_dan_sorot_dengan_escape():
    par = {3: "Nilai Rp 1.000 <lama> oleh PT Lama."}
    hasil = perubahan_html(par, [
        Usulan(3, "1.000", "2.000", "", "", ""),
        Usulan(3, "PT Lama", "PT A&B", "", "", ""),
    ])
    assert len(hasil) == 1 and hasil[0][0] == 3
    isi = hasil[0][1]
    assert "&lt;lama&gt;" in isi                       # teks asli di-escape
    assert '<del style="color:#b42318">1.000</del>' in isi
    assert '<mark style="background:#fff3a3">2.000</mark>' in isi
    assert '<mark style="background:#fff3a3">PT A&amp;B</mark>' in isi


def test_halaman_berubah_abaikan_spasi_dan_baris():
    halaman = ["Kop\r\nPendahuluan", "Nilai pasar Rp 595.400.000\r\n(LIMA RATUS\r\nRIBU)", "Lampiran"]
    assert halaman_berubah(halaman, ["595.400.000", "lima ratus ribu", "x"]) == [2]
    assert halaman_berubah(halaman, []) == []


@pytest.mark.skipif(not shutil.which("soffice"), reason="butuh LibreOffice")
def test_render_halaman_dari_pdf():
    from report_maker.pdf import docx_ke_pdf

    doc = Document()
    doc.add_paragraph("Halaman satu Rp 595.400.000")
    doc.add_page_break()
    doc.add_paragraph("Halaman dua")
    out = BytesIO()
    doc.save(out)
    pdf = docx_ke_pdf(out.getvalue())
    halaman = teks_per_halaman(pdf)
    assert len(halaman) == 2
    assert halaman_berubah(halaman, ["595.400.000"]) == [1]
    assert gambar_halaman(pdf, 2)[:8] == b"\x89PNG\r\n\x1a\n"
