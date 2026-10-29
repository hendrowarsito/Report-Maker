import json
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document

from report_maker.penyesuaian import (
    Usulan,
    daftar_paragraf,
    excel_ke_teks,
    terapkan_penyesuaian,
    usulkan_penyesuaian,
    validasi_usulan,
)


def _laporan() -> bytes:
    doc = Document()
    doc.add_paragraph("Objek berupa tanah SHM No. 0999 seluas 250 m².")
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.add_run("Nilai pasar Rp 1.000.000.000,00 ").bold = True
    p.add_run("(SATU MILIAR RUPIAH).")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Pemegang hak: Budi"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def test_daftar_paragraf_bernomor_tanpa_paragraf_kosong():
    par = daftar_paragraf(_laporan())
    assert list(par.values()) == [
        "Objek berupa tanah SHM No. 0999 seluas 250 m².",
        "Nilai pasar Rp 1.000.000.000,00 (SATU MILIAR RUPIAH).",
        "Pemegang hak: Budi",
    ]
    assert 1 not in par   # paragraf kosong dilewati, nomor tetap stabil


def test_excel_ke_teks_berlabel_sel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lahan"
    ws["A1"], ws["B1"] = "Luas tanah", 312
    ws["B3"] = 0.5
    wb.create_sheet("Kosong")
    out = BytesIO()
    wb.save(out)
    assert excel_ke_teks(out.getvalue()) == "## Sheet: Lahan\nA1: Luas tanah | B1: 312\nB3: 0,5"


def test_validasi_dan_terapkan_menjaga_format():
    src = _laporan()
    par = daftar_paragraf(src)
    u_ok = Usulan(0, "seluas 250 m²", "seluas 312 m²", "beda luas", "SHM.pdf hal. 1", "Luas 312 m2")
    u_nominal = Usulan(2, "1.000.000.000,00", "1.200.000.000,00", "", "", "")
    u_salah = Usulan(0, "tidak ada di paragraf", "x", "", "", "")
    u_nomor = Usulan(99, "a", "b", "", "", "")
    valid, tidak_valid = validasi_usulan(par, [u_ok, u_nominal, u_salah, u_nomor])
    assert valid == [u_ok, u_nominal]
    assert [a for _, a in tidak_valid] == [
        "Teks lama tidak ditemukan di paragraf", "Paragraf [99] tidak ada",
    ]

    hasil, jumlah = terapkan_penyesuaian(src, valid)
    assert jumlah == 2
    doc = Document(BytesIO(hasil))
    assert doc.paragraphs[0].text == "Objek berupa tanah SHM No. 0999 seluas 312 m²."
    run = doc.paragraphs[2].runs[0]
    assert run.text == "Nilai pasar Rp 1.200.000.000,00 " and run.bold


class _FakeClient:
    """Mencatat permintaan dan mengembalikan jawaban JSON tetap."""

    def __init__(self, jawaban: dict, stop_reason: str = "end_turn"):
        self.permintaan = None
        respons = SimpleNamespace(
            stop_reason=stop_reason,
            content=[SimpleNamespace(type="thinking"),
                     SimpleNamespace(type="text", text=json.dumps(jawaban))],
        )

        class _Stream:
            def __enter__(s): return s
            def __exit__(s, *a): return False
            def get_final_message(s): return respons

        def stream(**kw):
            self.permintaan = kw
            return _Stream()

        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=stream))


def test_usulkan_penyesuaian_bentuk_permintaan_dan_hasil():
    client = _FakeClient({
        "usulan": [{"paragraf": 0, "teks_lama": "250", "teks_baru": "312", "alasan": "a",
                    "sumber": "SHM.pdf hal. 1", "kutipan_sumber": "Luas 312"}],
        "catatan": "NIB tidak terbaca",
    })
    usulan, catatan = usulkan_penyesuaian(
        client, {0: "seluas 250 m²"}, [("SHM.pdf", b"%PDF-1.4")], [("LK.xlsx", "B1: 312")]
    )
    assert usulan == [Usulan(0, "250", "312", "a", "SHM.pdf hal. 1", "Luas 312")]
    assert catatan == "NIB tidak terbaca"

    req = client.permintaan
    assert req["model"] == "claude-opus-5-5"
    assert req["fallbacks"] == "default"
    assert req["output_config"]["format"]["type"] == "json_schema"
    pdf, excel, laporan = req["messages"][0]["content"]
    assert pdf["type"] == "document" and pdf["title"] == "SHM.pdf"
    assert excel["text"].startswith("# Lembar kerja Excel: LK.xlsx")
    assert excel["cache_control"] == {"type": "ephemeral"}
    assert "[0] seluas 250 m²" in laporan["text"]


@pytest.mark.parametrize("stop_reason", ["refusal", "max_tokens"])
def test_usulkan_penyesuaian_gagal_jelas(stop_reason):
    client = _FakeClient({"usulan": [], "catatan": ""}, stop_reason)
    with pytest.raises(RuntimeError):
        usulkan_penyesuaian(client, {0: "x"}, [], [("a.xlsx", "A1: 1")])
