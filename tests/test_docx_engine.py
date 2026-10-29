from io import BytesIO

from docx import Document

from report_maker.docx_engine import (
    isi_dokumen,
    nama_file_hasil,
    placeholders_dari_bytes,
)


def _docx(*paragraf_runs, header=None) -> bytes:
    """Buat .docx; tiap paragraf = list run, agar placeholder bisa dipecah."""
    doc = Document()
    for runs in paragraf_runs:
        p = doc.add_paragraph()
        for teks in runs:
            p.add_run(teks).bold = teks.startswith("**")
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    tabel = doc.add_table(rows=1, cols=1)
    tabel.cell(0, 0).text = "Sel {{nama}}"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _teks(docx_bytes: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(docx_bytes)).paragraphs]


def test_placeholder_terpecah_beberapa_run_dan_header_tabel():
    src = _docx(["Tanggal {{tgl_", "inspeksi}} oleh {{nama}}."], header="Kop {{nama}}")
    assert placeholders_dari_bytes(src) == {"tgl_inspeksi", "nama"}

    hasil, teks, sisa = isi_dokumen(src, {"tgl_inspeksi": "28 April 2026", "nama": "PT ABC"})
    assert _teks(hasil)[0] == "Tanggal 28 April 2026 oleh PT ABC."
    doc = Document(BytesIO(hasil))
    assert doc.sections[0].header.paragraphs[0].text == "Kop PT ABC"
    assert doc.tables[0].cell(0, 0).text == "Sel PT ABC"
    assert "Sel PT ABC" in teks
    assert sisa == []


def test_placeholder_tanpa_data_tetap_dan_dilaporkan():
    src = _docx(["{{ada}} dan {{tidak_ada}}"])
    hasil, _, sisa = isi_dokumen(src, {"ada": "1"})
    assert _teks(hasil)[0] == "1 dan {{tidak_ada}}"
    assert sisa == ["nama", "tidak_ada"]   # {{nama}} di tabel juga belum diisi


def test_find_replace_setelah_placeholder():
    src = _docx(["Pemberi tugas: {{nama}}"])
    hasil, _, _ = isi_dokumen(src, {"nama": "PT ABC"}, (("Pemberi tugas", "Klien"),))
    assert _teks(hasil)[0] == "Klien: PT ABC"


def test_format_run_tetap_terjaga():
    src = _docx(["**{{nama}}"])
    hasil, _, _ = isi_dokumen(src, {"nama": "PT ABC"})
    run = Document(BytesIO(hasil)).paragraphs[0].runs[0]
    assert run.text == "**PT ABC" and run.bold


def test_nama_file_hasil():
    data = {"nama": "PT ABC", "tgl": "28/04/2026"}
    assert nama_file_hasil("{{nama_file}}_terisi", "LP Final.docx", data) == "LP Final_terisi"
    assert nama_file_hasil("LP_{{nama}}_{{tgl}}", "a.docx", data) == "LP_PT ABC_28-04-2026"
    assert nama_file_hasil("LP_{{tidak_ada}}", "a.docx", data) == "LP"
    assert nama_file_hasil("{{tidak_ada}}", "a.docx", data) == "a"
