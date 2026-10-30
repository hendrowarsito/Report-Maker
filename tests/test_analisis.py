import json
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document

from report_maker.analisis import (
    Kandidat,
    Kemunculan,
    Temuan,
    analisis_dokumen,
    blok_sumber,
    rencana_perubahan,
    teks_pengganti,
)
from report_maker.penyesuaian import Usulan


@pytest.mark.parametrize("bentuk, teks_lama, nilai, hasil", [
    ("teks", "Jl. Lama No. 1", "Jl. Kalibata Raya No. 5", "Jl. Kalibata Raya No. 5"),
    ("angka", "1.000.000.000,00", "1.250.000.000", "1.250.000.000,00"),
    ("angka", "250", "312", "312"),
    ("angka", "1.000.000", "Rp 1.250.000.000", "1.250.000.000"),
    ("terbilang", "SATU MILIAR RUPIAH", "1.250.000.000",
     "SATU MILIAR DUA RATUS LIMA PULUH JUTA RUPIAH"),
    ("terbilang_judul", "Satu Miliar Rupiah", "1.250.000.000",
     "Satu Miliar Dua Ratus Lima Puluh Juta Rupiah"),
    ("angka", "250", "tiga ratus", None),
    # Tanpa pemisah ribuan di template: jangan tambahkan titik, nol di depan tetap
    ("angka", "0999", "1234", "1234"),
    ("angka", "250", "0312", "0312"),
    ("angka", "250", "1.312", "1.312"),
    # Placeholder tidak memberi contoh format: angka besar diberi titik ribuan
    ("angka", "{{nilai_pasar}}", "595400000", "595.400.000"),
    # Satuan & desimal dari sertifikat, gaya dari template
    ("angka", "455,00", "1.974,00 m²", "1.974,00"),
    ("angka", "455,00 m²", "1974", "1.974,00 m²"),
    ("angka", "250", "416 m2", "416"),
    ("angka", "1.000.000,00", "Rp 595.400.000", "595.400.000,00"),
    ("angka", "250", "312,5", "312,5"),
    ("terbilang", "X", "Rp 595.400.000", "LIMA RATUS SEMBILAN PULUH LIMA JUTA EMPAT RATUS RIBU RUPIAH"),
    # Nilai 0 (perhitungan belum selesai) tidak ditulis
    ("angka", "1.000.000,00", "0,00", None),
    ("teks", "{{nilai_pasar}}", "595.400.000", "595.400.000"),
])
def test_teks_pengganti(bentuk, teks_lama, nilai, hasil):
    assert teks_pengganti(Kemunculan(0, teks_lama, bentuk), nilai) == hasil


def test_rencana_perubahan_lewati_yang_sama_dan_bukan_angka():
    temuan = [
        Temuan("Nilai pasar", kemunculan=[
            Kemunculan(3, "1.000.000.000,00", "angka", "Bab 1.docx"),
            Kemunculan(3, "SATU MILIAR RUPIAH", "terbilang", "Bab 1.docx"),
            Kemunculan(0, "1.000.000.000,00", "angka", "Surat.docx"),
        ]),
        Temuan("Luas tanah", kemunculan=[Kemunculan(1, "312", "angka", "Bab 1.docx")]),
        Temuan("Alamat", kemunculan=[Kemunculan(0, "{{alamat}}", "teks", "Surat.docx")]),
    ]
    perubahan, peringatan = rencana_perubahan(
        temuan, {0: "1.250.000.000", 1: "312", 2: "Jl. X"}
    )
    # Satu pilihan nilai diterapkan seragam di semua file; luas 312 sudah sama
    assert perubahan == {
        "Bab 1.docx": [
            Usulan(3, "1.000.000.000,00", "1.250.000.000,00", "Nilai pasar", "", ""),
            Usulan(3, "SATU MILIAR RUPIAH", "SATU MILIAR DUA RATUS LIMA PULUH JUTA RUPIAH",
                   "Nilai pasar", "", ""),
        ],
        "Surat.docx": [
            Usulan(0, "1.000.000.000,00", "1.250.000.000,00", "Nilai pasar", "", ""),
            Usulan(0, "{{alamat}}", "Jl. X", "Alamat", "", ""),
        ],
    }
    assert peringatan == []

    _, peringatan = rencana_perubahan(temuan, {1: "tiga ratus"})
    assert peringatan == [
        "**Luas tanah**: “tiga ratus” bukan angka — 1 lokasi angka/terbilang tidak diubah "
        "(Bab 1 (1×)). Pilih sumber lain atau isi manual."
    ]


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LK Tanah"
    ws["B2"], ws["C2"] = "Luas tanah", 312
    ws["B3"], ws["C3"] = "Nilai pasar", 595400000
    ws["C3"].number_format = "General"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _docx(teks: str) -> bytes:
    d = Document()
    d.add_paragraph(teks)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def test_blok_sumber_per_jenis_file():
    pdf, sel = blok_sumber("SHM.pdf", b"%PDF-1.4")
    assert pdf["type"] == "document" and pdf["title"] == "SHM.pdf" and sel == {}
    xlsx, sel = blok_sumber("LK.xlsx", _xlsx())
    assert "B2: Luas tanah | C2: 312" in xlsx["text"] and sel == {
        "LK Tanah!B2": "Luas tanah", "LK Tanah!C2": "312",
        "LK Tanah!B3": "Nilai pasar", "LK Tanah!C3": "595.400.000",
    }
    docx, _ = blok_sumber("Proposal.docx", _docx("Pemberi tugas: PT ABC"))
    assert docx["text"] == "# Dokumen Word: Proposal.docx\n\nPemberi tugas: PT ABC"
    with pytest.raises(ValueError):
        blok_sumber("foto.jpg", b"")


def test_analisis_dokumen_nilai_lembar_kerja_dari_sel_asli():
    jawaban = {
        "temuan": [{
            "nama": "Luas tanah",
            "kandidat": [
                # AI salah menyalin angka lembar kerja → harus dikoreksi dari sel
                {"sumber": "LK.xlsx", "lokasi": "LK Tanah!C2", "nilai": "321"},
                # Angka sama dengan sel → tulisan berformat Claude dipertahankan
                {"sumber": "LK.xlsx", "lokasi": "LK Tanah!C3", "nilai": "595.400.000"},
                {"sumber": "SHM.pdf", "lokasi": "hal. 1", "nilai": "312"},
            ],
            "rekomendasi": "SHM.pdf",
            "konsisten": True,
            "penjelasan": "",
            "kemunculan": [{"dokumen": "Bab 2.docx", "paragraf": 0, "teks_lama": "250",
                            "bentuk": "angka"}],
        }],
        "catatan": "",
    }
    tercatat = {}
    respons = SimpleNamespace(
        stop_reason="end_turn",
        content=[SimpleNamespace(type="text", text=json.dumps(jawaban))],
    )

    class _Stream:
        def __enter__(s): return s
        def __exit__(s, *a): return False
        def get_final_message(s): return respons

    def stream(**kw):
        tercatat.update(kw)
        return _Stream()

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(stream=stream)))
    temuan, _ = analisis_dokumen(
        client,
        {"Surat.docx": {0: "Dengan hormat,"}, "Bab 2.docx": {0: "luas 250 m²"}},
        [("LK.xlsx", _xlsx()), ("SHM.pdf", b"%PDF-1.4")],
    )
    assert temuan[0].kandidat == [
        Kandidat("LK.xlsx", "LK Tanah!C2", "312", True, "lembar_kerja"),
        Kandidat("LK.xlsx", "LK Tanah!C3", "595.400.000", True, "lembar_kerja"),
        Kandidat("SHM.pdf", "hal. 1", "312", True, "sertifikat"),
    ]
    assert temuan[0].kemunculan == [Kemunculan(0, "250", "angka", "Bab 2.docx")]

    lk, shm, surat, bab2 = tercatat["messages"][0]["content"]
    assert shm["type"] == "document" and shm["cache_control"] == {"type": "ephemeral"}
    assert surat["text"] == "# Template: Surat.docx (paragraf bernomor)\n\n[0] Dengan hormat,"
    assert bab2["text"].startswith("# Template: Bab 2.docx") and "[0] luas 250 m²" in bab2["text"]


def test_rencana_perubahan_ikuti_gaya_huruf_template():
    temuan = [Temuan("Pemberi tugas", kemunculan=[
        Kemunculan(9, "PT JTRUST INVESTMENTS INDONESIA", "teks", "Surat.docx"),
        Kemunculan(23, "PT Lama", "teks", "Surat.docx"),
    ])]
    perubahan, _ = rencana_perubahan(temuan, {0: "PT JTrust Investments Indonesia"})
    # Penerima surat (kapital) sudah benar → tidak diubah; kemunculan lain diganti
    assert perubahan == {"Surat.docx": [
        Usulan(23, "PT Lama", "PT JTrust Investments Indonesia", "Pemberi tugas", "", ""),
    ]}


from report_maker.analisis import pilih_sumber, tebak_jenis_sumber  # noqa: E402


@pytest.mark.parametrize("nama, jenis", [
    ("SEQ 1 - SHM+05396+.pdf", "sertifikat"),
    ("Sertipikat HGB 123.pdf", "sertifikat"),
    ("49-JTII-Cost-Kasmiwati.xlsx", "lembar_kerja"),
    ("IMB No 45.pdf", "imb_pbg"),
    ("260916.001-SPN-A-JTrust.pdf", "proposal"),
    ("foto.pdf", "lainnya"),
])
def test_tebak_jenis_sumber(nama, jenis):
    assert tebak_jenis_sumber(nama) == jenis


def _t(kategori, *kandidat, bentuk="teks", rekomendasi=""):
    return Temuan("Data", list(kandidat), rekomendasi, True, "",
                  [Kemunculan(0, "x", bentuk, "a.docx")], kategori)


LK = "LK.xlsx"
SHM = "SHM 05396.pdf"


def test_data_sertifikat_dari_sertifikat_bila_jelas():
    t = _t("sertifikat",
           Kandidat(LK, "Laporan!B31", "1.974,00", True, "lembar_kerja"),
           Kandidat(SHM, "hal. 2", "1.974,00 m²", True, "sertifikat"),
           bentuk="angka", rekomendasi=LK)
    assert pilih_sumber(t) == (SHM, "Sertifikat")


def test_data_sertifikat_tidak_jelas_pakai_lembar_kerja():
    t = _t("sertifikat",
           Kandidat(SHM, "hal. 1", "Ujung Padang (dicoret)", False, "sertifikat"),
           Kandidat(LK, "Laporan!B15", "Desa Bandar Ratu", True, "lembar_kerja"))
    assert pilih_sumber(t) == (LK, "Lembar kerja (sertifikat tidak jelas)")
    t = _t("sertifikat", Kandidat(LK, "Laporan!B37", "Kasmiwati", True, "lembar_kerja"))
    assert pilih_sumber(t) == (LK, "Lembar kerja (sertifikat tidak memuat data ini)")


def test_nilai_nol_tidak_dipilih():
    t = _t("penilaian", Kandidat(LK, "Laporan!B67", "0,00", True, "lembar_kerja"), bentuk="angka")
    assert pilih_sumber(t)[0] == ""
    # Nilai teks "0" untuk data teks tetap boleh
    t = _t("penilaian", Kandidat(LK, "A1", "0", True, "lembar_kerja"))
    assert pilih_sumber(t)[0] == LK


def test_beberapa_sertifikat_pakai_yang_jelas_pertama():
    t = _t("sertifikat",
           Kandidat("SHM 1.pdf", "hal. 1", "05396", False, "sertifikat"),
           Kandidat("SHM 2.pdf", "hal. 1", "05379", True, "sertifikat"),
           Kandidat(LK, "B33", "05396", True, "lembar_kerja"))
    assert pilih_sumber(t) == ("SHM 2.pdf", "Sertifikat")
