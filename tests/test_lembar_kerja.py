import datetime
import json
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest

from report_maker.lembar_kerja import (
    Pemetaan,
    baca_sel,
    format_sel,
    kebutuhan_kunci,
    pemetaan_dari_json,
    pemetaan_ke_json,
    petakan_dengan_ai,
    sel_ke_teks,
    terapkan_pemetaan,
)


def _lembar_kerja() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LK Tanah"
    ws["B2"], ws["C2"] = "Alamat", "Jl. Kalibata Raya No. 5"
    ws["B3"], ws["C3"] = "Kelurahan", "Rawajati"
    ws["B4"], ws["C4"] = "Luas tanah (m²)", 312
    ws["B5"], ws["C5"] = "Tanggal penilaian", datetime.datetime(2026, 10, 1)
    rekap = wb.create_sheet("Rekap")
    rekap["A1"], rekap["B1"] = "Nilai pasar", 1_250_000_000
    rekap["B1"].number_format = "#,##0"
    rekap["A2"], rekap["B2"] = "Penyesuaian", 0.125
    rekap["B2"].number_format = "0.0%"
    rekap["A3"], rekap["B3"] = "Rumus tanpa nilai tersimpan", "=B1*2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize("nilai, fmt, teks", [
    (1_250_000_000, "#,##0", "1.250.000.000"),
    (1234.5, "#,##0.00", "1.234,50"),
    (0.125, "0.0%", "12,5%"),
    (0.25, "0%", "25%"),
    (312, "General", "312"),
    (0.5, "General", "0,5"),
    (datetime.datetime(2026, 10, 1), "dd/mm/yyyy", "1 Oktober 2026"),
    ("  teks ", "@", "teks"),
])
def test_format_sel(nilai, fmt, teks):
    assert format_sel(nilai, fmt) == teks


def test_baca_sel_dan_teks_untuk_ai():
    sel = baca_sel(_lembar_kerja())
    assert sel["LK Tanah!C2"] == "Jl. Kalibata Raya No. 5"
    assert sel["LK Tanah!C5"] == "1 Oktober 2026"
    assert sel["Rekap!B1"] == "1.250.000.000"
    assert sel["Rekap!B2"] == "12,5%"
    # Rumus yang belum pernah dihitung Excel tidak punya nilai tersimpan
    assert "Rekap!B3" not in sel
    teks = sel_ke_teks(sel)
    assert "## Sheet: LK Tanah\nB2: Alamat | C2: Jl. Kalibata Raya No. 5\n" in teks
    assert "## Sheet: Rekap\nA1: Nilai pasar | B1: 1.250.000.000" in teks


def test_kebutuhan_kunci_turunan_jadi_angka():
    assert kebutuhan_kunci({"alamat", "nilai_TERBILANG", "nilai", "luas_FORMAT"}) == {
        "alamat": "teks", "luas": "angka", "nilai": "angka",
    }


def test_terapkan_pemetaan_gabung_sel_dan_peringatan():
    sel = baca_sel(_lembar_kerja())
    data, peringatan = terapkan_pemetaan({
        "alamat": Pemetaan(["LK Tanah!C2", "LK Tanah!C3"], ", "),
        "nilai": Pemetaan(["Rekap!B1"]),
        "ragu": Pemetaan(["LK Tanah!C4"], keyakinan="rendah"),
        "hilang": Pemetaan([], keyakinan="tidak_ditemukan"),
        "salah_sel": Pemetaan(["LK Tanah!Z99"]),
    }, sel)
    assert data == {
        "alamat": "Jl. Kalibata Raya No. 5, Rawajati",
        "nilai": "1.250.000.000",
        "ragu": "312",
    }
    assert peringatan == {
        "ragu": "Keyakinan rendah — cek LK Tanah!C4",
        "hilang": "Tidak ditemukan di lembar kerja",
        "salah_sel": "Sel kosong/tidak ada: LK Tanah!Z99",
    }


def test_simpan_dan_muat_pemetaan():
    peta = {"alamat": Pemetaan(["LK Tanah!C2", "LK Tanah!C3"], ", ", "tinggi", "label Alamat")}
    assert pemetaan_dari_json(pemetaan_ke_json(peta)) == peta
    with pytest.raises(ValueError):
        pemetaan_dari_json("{bukan json")
    with pytest.raises(ValueError):
        pemetaan_dari_json('{"tanpa": "pemetaan"}')


def test_petakan_dengan_ai_bentuk_permintaan():
    tercatat = {}
    jawaban = {
        "pemetaan": [
            {"key": "alamat", "sel": ["LK Tanah!C2"], "pemisah": " ",
             "keyakinan": "tinggi", "alasan": "label Alamat di B2"},
            {"key": "tidak_diminta", "sel": [], "pemisah": " ",
             "keyakinan": "tidak_ditemukan", "alasan": ""},
        ],
        "catatan": "",
    }
    respons = SimpleNamespace(
        stop_reason="end_turn",
        content=[SimpleNamespace(type="text", text=json.dumps(jawaban))],
    )

    class _Stream:
        def __enter__(s): return s
        def __exit__(s, *a): return False
        def get_final_message(s): return respons

    def stream(**kw):
        tercatat.update(kw)
        return _Stream()

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(stream=stream)))
    peta, _ = petakan_dengan_ai(
        client, {"LK Tanah!C2": "Jl. X"},
        {"alamat": "teks", "nilai": "angka"},
        {"alamat": "terletak di {{alamat}}", "nilai_TERBILANG": "sebesar ({{nilai_TERBILANG}})"},
    )
    assert peta == {"alamat": Pemetaan(["LK Tanah!C2"], " ", "tinggi", "label Alamat di B2")}
    lembar, daftar = tercatat["messages"][0]["content"]
    assert "C2: Jl. X" in lembar["text"] and lembar["cache_control"]
    # Konteks key sumber diambil dari placeholder turunannya
    assert '`nilai` (angka) — konteks: "sebesar ({{nilai_TERBILANG}})"' in daftar["text"]
